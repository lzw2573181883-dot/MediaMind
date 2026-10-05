# 文件路径：backend/app/agent/tools.py
# 职责：定义供大模型自主决策调用的两个专精工具（RAG 时间戳检索 + Word 文档导出）
import os
from typing import Dict, Any
from docx import Document

# 直接复用 app.rag 全局唯一单例，杜绝重复创建
from app.rag import hybrid_retriever, reranker_service

# ==========================================
# 道具 1：给大模型看的【工具箱说明书】(JSON Schema)
# ==========================================
TOOLS_SCHEMA = [
    {
        "type": "function",
        "function": {
            "name": "search_video_rag",
            "description": "【核心知识检索工具】：当用户询问视频中的具体内容、参数、事实、优缺点或任何细节时必须调用此工具。通过双路混合检索与交叉重排，从原片中精选出最匹配的台词切片与精准秒级时间戳。",
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "针对视频细节提取出的搜索关键词或查询短句，例如：'4K 120帧录制发热情况与续航时长'"
                    },
                    "bvid": {
                        "type": "string",
                        "description": "视频的唯一标识 BVID，如 BV1cSec6tEux"
                    }
                },
                "required": ["query"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "export_word_document",
            "description": "【Word文档导出工具】：当用户明确表达想'导出笔记'、'下载Word'、'保存为文档'、'下载复习资料'时调用此工具。将总结内容整理并排版保存为标准的 Microsoft Word (.docx) 办公文档。",
            "parameters": {
                "type": "object",
                "properties": {
                    "filename": {
                        "type": "string",
                        "description": "要保存的 Word 文档名称，例如：'相机评测学习笔记.docx'"
                    },
                    "title": {
                        "type": "string",
                        "description": "Word 文档内的一级大标题，例如：'4K高规格相机全面评测核心笔记'"
                    },
                    "content": {
                        "type": "string",
                        "description": "要写入 Word 文档的完整正文内容（包含各章节核心要点与时间戳信息）"
                    }
                },
                "required": ["filename", "title", "content"]
            }
        }
    }
]


# ==========================================
# 道具 2：真正干活的 Python 函数（执行体）
# ==========================================
class AgentTools:
    """
    真正干活的工具执行中心：共享全局 RAG 检索器，开箱即用
    """
    def __init__(self, export_dir: str = "./data/exports"):
        # 共享全局单例，零冗余
        self.retriever = hybrid_retriever
        self.reranker = reranker_service
        self.export_dir = export_dir
        # 自动在 backend 下创建 data/exports 文件夹，专门存放导出的 Word
        os.makedirs(self.export_dir, exist_ok=True)

    async def search_video_rag(self, query: str, bvid: str = "") -> Dict[str, Any]:
        """
        工具 1：调用双路检索 + 交叉重排，返回格式化上下文与黄金切片
        """
        candidates = await self.retriever.search(query=query, bvid=bvid, top_k=15)
        if not candidates:
            return {
                "status": "empty",
                "message": "未在视频中检索到相关切片",
                "context_text": "未检索到相关内容，请如实告知用户相关知识点原片未提及。",
                "chunks": []
            }

        golden_chunks = await self.reranker.rerank(query=query, candidates=candidates, top_k=3)
        
        # 将 Top-3 黄金切片拼装成带时间戳锚点的文本上下文，供大模型阅读
        context_blocks = []
        for idx, c in enumerate(golden_chunks, 1):
            time_tag = c.get("time_str", "00:00")
            snippet = c.get("text", "")
            context_blocks.append(f"【参考切片 {idx} (时间戳 [{time_tag}])】\n{snippet}")
        context_text = "\n\n".join(context_blocks)

        return {
            "status": "success",
            "matched_count": len(golden_chunks),
            "context_text": context_text,
            "chunks": golden_chunks
        }

    def export_word_document(self, filename: str, title: str, content: str) -> Dict[str, Any]:
        """
        工具 2：调用 python-docx，生成真正的 .docx 办公文档
        """
        if not filename.endswith(".docx"):
            filename += ".docx"

        # 拿出空白 Word 文档并排版
        doc = Document()
        doc.add_heading(title, level=1)

        for line in content.split("\n"):
            stripped_line = line.strip()
            if stripped_line:
                doc.add_paragraph(stripped_line)

        # 保存到本地磁盘 ./data/exports/xxx.docx（防御性处理：若文件被Word等软件打开占用，自动添加时间戳另存，绝不报错崩盘）
        filepath = os.path.join(self.export_dir, filename)
        try:
            doc.save(filepath)
        except PermissionError:
            import time
            base, ext = os.path.splitext(filename)
            filename = f"{base}_{int(time.time())}{ext}"
            filepath = os.path.join(self.export_dir, filename)
            doc.save(filepath)

        return {
            "status": "success",
            "message": f"🎉 专属 Word 笔记已成功生成并保存！",
            "filename": filename,
            "filepath": os.path.abspath(filepath)
        }

# 实例化全局单例工具集
agent_tools = AgentTools()
